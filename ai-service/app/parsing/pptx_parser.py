"""PPTX 파싱 — 슬라이드별 텍스트 추출.

python-pptx로 구현. "주장(claim)" 단위 추출은 여기서 하지 않는다 — 슬라이드
원문 텍스트를 통째로 넘기면 app/llm/gateway_client.py의 LLM 프롬프트가 알아서
핵심 주장을 뽑아내도록 이미 설계되어 있어서(시스템 프롬프트 1번 항목 참고),
여기는 슬라이드당 텍스트 하나만 만들면 된다.

- slide_index는 0부터 시작 (python-pptx 관례, API_명세서 §2.3.1 참고 —
  프론트가 사람이 보는 번호로 표시할 땐 +1 하므로 여기선 그대로 0-based 유지).
- 텍스트 상자/제목/본문/표를 전부 훑고, 그룹 도형 안의 텍스트도 재귀적으로 뽑는다.
- 도형들을 화면 위치(top, left) 기준으로 정렬해서 대략적인 읽는 순서를 맞춘다
  (pptx XML 순서는 z-order/생성 순서라 실제 읽는 순서랑 다를 수 있음 — 완벽한
  정렬은 아니지만 이 정도면 충분).
- 발표자 노트(speaker notes)는 포함하지 않는다 — 청중이 보는 슬라이드 내용과
  실제 발화를 비교하는 게 목적이라, 발표자만 보는 비공개 메모는 제외.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError


@dataclass
class Slide:
    slide_index: int  # 0-based
    text: str  # 슬라이드 내 텍스트 원문(제목+본문+표 등, 위치순으로 이어붙임)


class PptxParseError(Exception):
    """pptx 파싱 실패. app.schemas.job.ErrorCode.PPTX_PARSE_FAILED에 대응."""


def parse_pptx(pptx_path: Path) -> list[Slide]:
    """pptx 파일을 슬라이드 리스트로 파싱한다.

    Raises:
        PptxParseError: 파일이 없거나, pptx 형식이 아니거나, 슬라이드가
            하나도 없는 경우.
    """
    if not pptx_path.exists():
        raise PptxParseError(f"PPTX 파일을 찾을 수 없음: {pptx_path}")

    try:
        presentation = Presentation(str(pptx_path))
    except PackageNotFoundError as exc:
        raise PptxParseError(f"유효한 PPTX 파일이 아님: {pptx_path}") from exc
    except Exception as exc:  # python-pptx가 다양한 예외를 던질 수 있어 통일
        raise PptxParseError(f"PPTX 파싱 실패: {exc}") from exc

    slides = [
        Slide(slide_index=i, text=_extract_slide_text(slide))
        for i, slide in enumerate(presentation.slides)
    ]

    if not slides:
        raise PptxParseError(f"슬라이드가 하나도 없음: {pptx_path}")

    return slides


def _extract_slide_text(slide) -> str:
    """슬라이드 하나의 모든 도형에서 텍스트를 뽑아 위치순으로 이어붙인다."""
    parts = _shape_texts(slide.shapes)
    return "\n".join(parts)


def _shape_texts(shapes) -> list[str]:
    """도형 목록(대략 위치순 정렬)에서 텍스트를 재귀적으로 뽑는다.

    그룹 도형(MSO_SHAPE_TYPE.GROUP)은 내부 도형들을 재귀적으로 훑는다.
    """
    ordered = sorted(shapes, key=lambda s: (_shape_pos(s, "top"), _shape_pos(s, "left")))

    texts: list[str] = []
    for shape in ordered:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_shape_texts(shape.shapes))
            continue

        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)
        elif shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))

    return texts


def _shape_pos(shape, attr: str) -> int:
    """shape.top / shape.left. 위치 정보가 없는(None) 도형은 맨 앞(0)으로 취급."""
    value = getattr(shape, attr, None)
    return int(value) if value is not None else 0
