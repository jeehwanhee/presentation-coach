"""app/parsing/pptx_parser.py 단위 테스트.

python-pptx로 실제 .pptx 파일을 즉석에서 만들어서 그걸 다시 파싱하는 방식으로
검증한다(파일 포맷 자체를 왕복 검증하는 셈이라 fixture 파일을 따로 안 둬도 됨).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.parsing import pptx_parser as pp


def _blank_layout(prs: Presentation):
    # 인덱스 6은 python-pptx 기본 템플릿의 "빈 화면" 레이아웃.
    return prs.slide_layouts[6]


def _add_textbox(slide, text: str, top_in: float, left_in: float = 0.5):
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(4), Inches(1))
    box.text_frame.text = text
    return box


class TestParsePptxBasic:
    def test_missing_file_raises(self, tmp_path: Path):
        with pytest.raises(pp.PptxParseError, match="찾을 수 없음"):
            pp.parse_pptx(tmp_path / "missing.pptx")

    def test_non_pptx_file_raises(self, tmp_path: Path):
        bad = tmp_path / "not_a_pptx.pptx"
        bad.write_text("이건 그냥 텍스트 파일임")

        with pytest.raises(pp.PptxParseError):
            pp.parse_pptx(bad)

    def test_empty_presentation_raises(self, tmp_path: Path):
        prs = Presentation()  # 기본 템플릿, 슬라이드 0개
        path = tmp_path / "empty.pptx"
        prs.save(str(path))

        with pytest.raises(pp.PptxParseError, match="슬라이드가 하나도 없음"):
            pp.parse_pptx(path)

    def test_single_slide_text_extracted(self, tmp_path: Path):
        prs = Presentation()
        slide = prs.slides.add_slide(_blank_layout(prs))
        _add_textbox(slide, "안녕하세요 발표 시작합니다", top_in=0.5)
        path = tmp_path / "one_slide.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert len(slides) == 1
        assert slides[0].slide_index == 0
        assert "안녕하세요 발표 시작합니다" in slides[0].text

    def test_multiple_slides_index_0_based_in_order(self, tmp_path: Path):
        prs = Presentation()
        for i in range(3):
            slide = prs.slides.add_slide(_blank_layout(prs))
            _add_textbox(slide, f"슬라이드 {i}", top_in=0.5)
        path = tmp_path / "three_slides.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert [s.slide_index for s in slides] == [0, 1, 2]
        assert [s.text for s in slides] == ["슬라이드 0", "슬라이드 1", "슬라이드 2"]


class TestShapeExtraction:
    def test_multiple_textboxes_ordered_by_position(self, tmp_path: Path):
        prs = Presentation()
        slide = prs.slides.add_slide(_blank_layout(prs))
        # 일부러 아래쪽 텍스트를 먼저 추가 — 도형 추가 순서가 아니라 위치(top)로
        # 정렬되는지 확인하려는 목적.
        _add_textbox(slide, "아래 텍스트", top_in=3.0)
        _add_textbox(slide, "위 텍스트", top_in=0.5)
        path = tmp_path / "ordered.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert slides[0].text == "위 텍스트\n아래 텍스트"

    def test_table_text_extracted_with_separator(self, tmp_path: Path):
        prs = Presentation()
        slide = prs.slides.add_slide(_blank_layout(prs))
        table_shape = slide.shapes.add_table(
            rows=2, cols=2, left=Inches(0.5), top=Inches(0.5), width=Inches(4), height=Inches(1.5)
        )
        table = table_shape.table
        table.cell(0, 0).text = "항목"
        table.cell(0, 1).text = "값"
        table.cell(1, 0).text = "이탈률"
        table.cell(1, 1).text = "30%"
        path = tmp_path / "table.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert "항목 | 값" in slides[0].text
        assert "이탈률 | 30%" in slides[0].text

    def test_group_shape_text_extracted_recursively(self, tmp_path: Path):
        prs = Presentation()
        slide = prs.slides.add_slide(_blank_layout(prs))
        box1 = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(1))
        box1.text_frame.text = "그룹 안 텍스트 1"
        box2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2), Inches(1))
        box2.text_frame.text = "그룹 안 텍스트 2"
        group = slide.shapes.add_group_shape([box1, box2])
        # add_group_shape로 묶은 뒤에도 개별 shape가 최상위 shapes에 남지 않고
        # group 안으로 들어가는지는 python-pptx 버전에 따라 다를 수 있어서,
        # 그룹 텍스트가 최소 한 번은 결과에 포함되는지만 확인(중복 카운트 방지
        # 목적의 엄격한 assert는 피함).
        path = tmp_path / "group.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert "그룹 안 텍스트 1" in slides[0].text
        assert "그룹 안 텍스트 2" in slides[0].text

    def test_empty_slide_yields_empty_text_not_error(self, tmp_path: Path):
        prs = Presentation()
        prs.slides.add_slide(_blank_layout(prs))  # 아무 도형도 안 넣음
        path = tmp_path / "blank_slide.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert len(slides) == 1
        assert slides[0].text == ""

    def test_speaker_notes_are_not_included(self, tmp_path: Path):
        prs = Presentation()
        slide = prs.slides.add_slide(_blank_layout(prs))
        _add_textbox(slide, "슬라이드 본문", top_in=0.5)
        slide.notes_slide.notes_text_frame.text = "이건 발표자만 보는 비공개 메모"
        path = tmp_path / "with_notes.pptx"
        prs.save(str(path))

        slides = pp.parse_pptx(path)

        assert "슬라이드 본문" in slides[0].text
        assert "비공개 메모" not in slides[0].text
